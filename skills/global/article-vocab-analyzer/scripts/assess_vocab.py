# Global/$name: SKILL.md 脚本
# -*- coding: utf-8 -*-
"""对任意英文文章做四词表（KET/PET/NGSL/NAWL）词汇分析 + 可选 CEFR 难度标注。

判定流水线（优先级从高到低）:
  1. 数字 -> num
  2. 专名（首字母大写且 lemma 未被词表收录）-> proper
  3. 多词短语整查（最长 5->2 优先）-> 命中最低层词表
  4. 单词：lemma 候选归一 + 研究版屈折词形兜底，逐词表命中
  5. 未命中 -> 超纲 out
分层归属：一个 token 命中多个词表时按 KET<PET<NGSL<NAWL 取最低层只归一层，覆盖率闭环不重叠。

用法:
  python assess_vocab.py --text "全文..."
  python assess_vocab.py --file 课文.txt
  python assess_vocab.py --xlsx 教学.xlsx --sheet "C4 句库时间轴" --col 2
  python assess_vocab.py --file 课文.txt --cefr cefr.json   # 附加 CEFR 难度标注

输出(写到 cwd):
  vocab_report.xlsx  覆盖率(各层+合并+超纲+专名/数字+每词表命中明细+超纲词)
  vocab_report.json  结构化结果(含 CEFR 等级/归属层)
词表: 读 <skill>/wordlists/registry.json 声明的顺序；把任意词表(含未来中高考)做同格式 json 放入即并行纳入。
词表格式: {"entries":[{"head","pos","raw"}],"forms":[...可选全屈折词形]}
"""
import os, sys, re, json, argparse, collections, glob, math


# 词汇库(SQLite累积/Excel导出)按需导入，透明地在读词表脚本里服务 --db。
def _load_word_bank():
    import importlib.util
    spec = importlib.util.spec_from_file_location("word_bank",
        os.path.join(os.path.dirname(os.path.abspath(__file__)), "word_bank.py"))
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
    return m

# ---------------- 词形归一（沿用旧技能踩坑结论） ----------------
IRREG = {
    "am": "be", "is": "be", "are": "be", "was": "be", "were": "be", "been": "be", "being": "be",
    "has": "have", "had": "have", "having": "have",
    "does": "do", "did": "do", "done": "do", "doing": "do",
    "says": "say", "said": "say", "went": "go", "gone": "go", "goes": "go",
    "saw": "see", "seen": "see", "made": "make", "making": "make",
    "took": "take", "taken": "take", "taking": "take",
    "got": "get", "gotten": "get", "gave": "give", "given": "give",
    "came": "come", "found": "find", "left": "leave", "thought": "think", "thoughts": "thought",
    "bought": "buy", "brought": "bring", "told": "tell",
    "felt": "feel", "kept": "keep", "knew": "know", "known": "know",
    "put": "put", "ran": "run", "sat": "sit", "stood": "stand", "wrote": "write",
    "written": "write", "won": "win", "wore": "wear", "children": "child",
    "people": "person", "men": "man", "women": "woman", "feet": "foot",
    "teeth": "tooth", "mice": "mouse", "best": "good", "better": "good",
    "worst": "bad", "worse": "bad", "babies": "baby", "watched": "watch",
}
CONTR = {"can't": "can", "won't": "will", "shan't": "shall",
         "don't": "do", "doesn't": "do", "didn't": "do",
         "isn't": "be", "aren't": "be", "wasn't": "be", "weren't": "be",
         "hasn't": "have", "haven't": "have", "hadn't": "have",
         "couldn't": "can", "shouldn't": "should", "wouldn't": "will",
         "mustn't": "must", "mightn't": "might",
         "i'm": "be", "you're": "be", "we're": "be", "they're": "be",
         "i've": "have", "you've": "have", "we've": "have", "they've": "have",
         "i'll": "will", "you'll": "will", "he'll": "will", "she'll": "will",
         "they'll": "will", "we'll": "will", "it's": "it", "there's": "be"}
# 以 -er 结尾的演员性名词：剥离会误判为比较级(presenter→present)。
AGENT_ER_NOUNS = {
    'presenter', 'teacher', 'driver', 'player', 'winner', 'worker',
    'writer', 'reader', 'speaker', 'listener', 'viewer', 'employer',
    'manager', 'owner', 'maker', 'seller', 'builder', 'thinker',
}
NUMBERS = {
    "zero", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine", "ten",
    "eleven", "twelve", "thirteen", "fourteen", "fifteen", "sixteen", "seventeen",
    "eighteen", "nineteen", "twenty", "thirty", "forty", "fifty", "sixty", "seventy",
    "eighty", "ninety", "hundred", "thousand", "million", "billion",
}
TOK_RE = re.compile(r"[A-Za-z']+|\d+")

def tokenize(s):
    # 剔除纯引号 token（单独成词的撇号），避免 `'` 被当词计入超纲
    return [t for t in TOK_RE.findall(s) if t not in ("'", chr(0x2019), chr(0x2018))]

def base_word(tok):
    # 版式引号(弯/全角撇号)归一到 ASCII，避免 `'Can`、`'` 被当 token
    t = tok.lower()
    t = t.replace('‘', "'").replace('’', "'").replace('“', '').replace('”', '')
    if t.startswith("'") and not t.startswith("'s"):
        t = t[1:]
    if t.endswith("'") and len(t) > 1:
        t = t[:-1]
    if t in CONTR: return CONTR[t]
    if t.endswith("'s"): t = t[:-2]
    elif t.endswith("s'"): t = t[:-2]
    elif t.endswith("n't"): t = t[:-2]
    elif "'" in t: t = t.split("'", 1)[0]
    return t

def is_cap_word(tok):
    """专名判定用：跳过前导引号/撇号后，看第一个字母是否大写。避免 `'Mo'`(带引号昵称)错判小写→超纲。"""
    t = tok
    while t and t[0] in "'‘’“”\"`":
        t = t[1:]
    if not t:
        return False
    return t[0].isupper()

def strip_quotes(tok):
    """剥离专名两侧成对引号/撇号，返回干净显示形式（如 'Mo' → Mo）。"""
    t = tok
    while t and t[0] in "'‘’“”\"`":
        t = t[1:]
    while t and t[-1] in "'‘’“”`":
        t = t[:-1]
    return t

def stem_suffix(w):
    if w.endswith("ies") and len(w) > 3: return w[:-3] + "y"
    if w.endswith("es") and len(w) > 2:
        s = w[:-2]
        return s if s[-1] in "sxz" or s.endswith("ch") or s.endswith("sh") else s + "e"
    if w.endswith("s") and not w.endswith("ss") and len(w) > 2: return w[:-1]
    if w.endswith("ing"):
        b = w[:-3]
        cands = [b, b + "e"]
        if len(b) >= 2 and b[-1] == b[-2]: cands.append(b[:-1])
        for c in cands:
            if c and re.fullmatch(r"[a-z]+", c): return c
        return b
    if w.endswith("ed"):
        b = w[:-2]
        if len(b) >= 2 and b[-1] == b[-2]: b = b[:-1]
        for c in (b, b + "e"):
            if re.fullmatch(r"[a-z]+", c): return c
        return b
    if w.endswith("er") and w not in AGENT_ER_NOUNS: return w[:-2]
    if w.endswith("est"): return (w[:-3] or w)
    return w

def lemma_candidates(word):
    t = base_word(word)
    if not t: return []
    out = [t]
    if t in IRREG: out.append(IRREG[t])
    s = stem_suffix(t)
    if s != t: out.extend([s, IRREG.get(s, s)])
    if t.endswith("ing"):
        b = t[:-3]
        if len(b) >= 2 and b[-1] == b[-2]: out.extend([b + "e", b[:-1]])
        out.append(b)
    if t.endswith(("es", "s")) and len(t) > 2 and not t.endswith("ing"):
        out.append(t[:-1])
    return list(dict.fromkeys(c for c in out if re.fullmatch(r"[a-z]+", c)))

def display_word(form, lemma):
    """卡面显示口径（2026-10-13 定案）：常规屈折 → 词原；不规则词形 → 保留表层形式。
    dancers→dancer、babies→baby、goes→go、taking→take；children/feet/went/got/came/
    best/better/made/said/left 等保留原形式，不折叠成 child/foot。
    判定：lemma 可由常规后缀规则还原 → 词原；只能经 IRREG 解释 → 保留原形式。
    （层级/CEFR 仍按词族原形标注，调用方自行取 lemma 的层。）"""
    base = base_word(form).lower()
    lemma = (lemma or "").lower()
    if base == lemma or stem_suffix(base) == lemma:
        return lemma                                   # 本形即词原，或常规后缀可还原
    regs = set()                                       # 常规屈折全备选（补 stem_suffix 拼写缺陷）
    if base.endswith("ies") and len(base) > 3: regs.add(base[:-3] + "y")
    if base.endswith("es"):
        regs.add(base[:-2]); regs.add(base[:-1]); regs.add(base[:-2] + "e")
    if base.endswith("s"): regs.add(base[:-1])
    if base.endswith("ing"):
        b = base[:-3]; regs |= {b, b + "e"}
        if len(b) >= 2 and b[-1] == b[-2]: regs.add(b[:-1])
    if base.endswith("ed"):
        b = base[:-2]; regs |= {b, b + "e"}
        if len(b) >= 2 and b[-1] == b[-2]: regs.add(b[:-1])
    if base.endswith("er") and base not in AGENT_ER_NOUNS: regs.add(base[:-2])
    if base.endswith("est"): regs.add(base[:-3])
    if lemma in regs:
        return lemma                                   # 规则屈折可还原 → 词原
    if IRREG.get(base) == lemma or IRREG.get(stem_suffix(base)) == lemma:
        return base                                    # 不规则词形：保留原形式
    return lemma                                       # 兜底：非常规拼写仍用词原

def split_head(head):
    parts = []
    for seg in head.lower().split("/"):
        seg = seg.strip()
        if seg: parts.append(seg)
    return parts or [head.lower().strip()]

# ---------------- 词表加载 ----------------
def load_wordlist(path):
    d = json.load(open(path, encoding="utf-8"))
    words, phrs, pos, forms = set(), {}, {}, set()
    for e in d.get("entries", []):
        for head in split_head(e["head"]):
            if " " in head:
                phrs.setdefault(head, e["pos"])
            else:
                words.add(head); pos.setdefault(head, e["pos"])
    forms.update(p.lower() for p in d.get("forms", []))
    return {"words": words, "phrs": phrs, "pos": pos, "forms": forms,
            "src": os.path.basename(path)}

def discover_wordlists(skill_dir):
    wd = os.path.join(skill_dir, "wordlists")
    reg = os.path.join(wd, "registry.json")
    order = json.load(open(reg, encoding="utf-8"))["wordlists"] if os.path.exists(reg) else []
    order = [w["key"] for w in order]
    found = {os.path.splitext(os.path.basename(jf))[0]: jf
             for jf in glob.glob(os.path.join(wd, "*.json"))
             if os.path.basename(jf) not in ("registry.json", "cefr.json", "phrase_extra.json")}
    keys = [k for k in order if k in found] + sorted(set(found) - set(order))
    wls = [(k, found[k], load_wordlist(found[k])) for k in keys]
    # 教学扩展短语种子库 phrase_extra.json: 官方未收录的固定结构(如 have got=have)。
    # 把其短语并入 "layer" 指定的词表 phrs, 使分词链零改动整体命中; 未指定 layer 则并入最低层。
    xp = os.path.join(wd, "phrase_extra.json")
    if os.path.exists(xp):
        try:
            xd = json.load(open(xp, encoding="utf-8"))
            by_layer = {}
            for e in xd.get("entries", []):
                if " " not in e["head"] and " " not in e.get("raw", ""):
                    continue
                head = e["head"].lower().strip()
                # extra 级默认并入最低层 KET; 条目可显式指定 layer 覆盖
                layer = (e.get("layer") or "KET").upper()
                by_layer.setdefault(layer, {})[head] = e["pos"]
            for ki, (k, p, wl) in enumerate(wls):
                if k in by_layer:
                    wl["phrs"].update(by_layer[k])
        except Exception as ex:
            sys.stderr.write("phrase_extra.json 解析失败, 忽略: %s\n" % ex)
    return wls

# ---------------- CEFR 标注 ----------------
def load_cefr(path):
    """cefr.json: {"words": {"head": "A2", ...}}. 无则返回 None。"""
    if not path or not os.path.exists(path):
        return None
    return json.load(open(path, encoding="utf-8")).get("words", {})

# 高频不规则屈折/变体 -> 基础形(cefr.json 收录的规范 head)。
# 仅用于 CEFR 标注宽容回退(见 cefr_level), 不触碰全局分词/分层/词库主键的 lemma。
_CEFR_FOLD = {
    "an": "a",
    "chose": "choose", "chosen": "choose",
    "decided": "decide",
    "healthier": "healthy", "healthiest": "healthy",
    "flew": "fly", "flown": "fly", "gone": "go",
    "took": "take", "taken": "take", "left": "leave",
    "ran": "run", "began": "begin", "became": "become",
    "brought": "bring", "built": "build", "caught": "catch",
    "came": "come", "drank": "drink", "drove": "drive",
    "ate": "eat", "found": "find", "got": "get", "gave": "give",
    "grew": "grow", "held": "hold", "kept": "keep", "knew": "know",
    "made": "make", "paid": "pay", "sat": "sit", "sang": "sing",
    "sold": "sell", "sent": "send", "slept": "sleep",
    "spent": "spend", "stood": "stand", "taught": "teach",
    "thought": "think", "understood": "understand", "won": "win",
    "wore": "wear", "meant": "mean", "broke": "break", "broken": "break",
    "spoke": "speak", "spoken": "speak", "wrote": "write", "written": "write",
}

def cefr_level(cefr, lemma):
    if not cefr: return ""
    v = cefr.get(lemma)
    if v: return v
    # 宽容归一: lemma 直查未命中时, 先用词形候选折到基础形, 再查显式不规则折叠表。
    # 只影响 CEFR 标注列, 不触碰分词/分层归属/词库主键的 lemma。
    for c in lemma_candidates(lemma):
        v = cefr.get(c)
        if v: return v
    return cefr.get(_CEFR_FOLD.get(lemma, ""), "")

# ---------------- 课文加载 ----------------
def load_sentences(args):
    if args.text:
        text = args.text
    elif args.stdin:
        text = sys.stdin.read()
    elif args.xlsx:
        import openpyxl
        wb = openpyxl.load_workbook(args.xlsx, read_only=True)
        ws = wb[args.sheet]
        return [str(r[args.col]).strip() for r in ws.iter_rows(min_row=2, values_only=True)
                if r[args.col] and str(r[args.col]).strip()]
    else:
        if not args.file: sys.exit("需提供 --text/--file/--xlsx/--stdin 之一")
        text = open(args.file, encoding="utf-8").read()
    sents = []
    for para in re.split(r"\n\s*\n", text):
        for s in re.split(r"(?<=[.!?])\s+", para.strip()):
            if s.strip(): sents.append(s.strip())
    return sents

# ---------------- 主判定 ----------------
def assess(sentences, wls):
    n = len(wls)
    per_key = [{} for _ in range(n)]
    hit_tok = [set() for _ in range(n)]
    out_items = []
    numbers_n = proper_n = 0
    for si, s in enumerate(sentences):
        toks = tokenize(s)
        n_l = len(toks)
        for ki, (key, path, wl) in enumerate(wls):
            used = [False] * n_l
            for L in range(5, 1, -1):
                i = 0
                while i + L <= n_l:
                    if any(used[i + k] for k in range(L)):
                        i += 1; continue
                    seq = " ".join(t.lower() for t in toks[i:i + L])
                    if seq in wl["phrs"]:
                        a = per_key[ki].setdefault(("phrase", seq),
                            {"type": "phrase", "lemma": seq, "form": seq, "pos": wl["phrs"][seq],
                             "freq": 0, "nsent": set(), "examples": []})
                        a["freq"] += 1; a["nsent"].add(si)
                        if len(a["examples"]) < 2: a["examples"].append(s)
                        for k in range(L):
                            used[i + k] = True; hit_tok[ki].add((si, i + k))
                    i += 1
        for ti, t in enumerate(toks):
            if re.fullmatch(r"\d+", t) or base_word(t) in NUMBERS:
                numbers_n += 1; out_items.append({"kind": "num", "form": t, "lemma": base_word(t), "hit": []}); continue
            cands = lemma_candidates(t)
            bw = base_word(t).lower()
            hits = [ki for ki, (k, p, wl) in enumerate(wls)
                    if (bw in wl["forms"]) or any(c in wl["words"] for c in cands)]
            is_cap = is_cap_word(t)
            if not hits:
                fb = cands[0] if cands else (bw or t.lower())
                if is_cap:
                    proper_n += 1; out_items.append({"kind": "proper", "form": strip_quotes(t), "lemma": fb, "hit": []})
                else:
                    out_items.append({"kind": "out", "form": t, "lemma": fb, "hit": [], "sent": s})
            else:
                lm = next((c for c in cands if c in wls[hits[0]][2]["words"]), cands[0])
                out_items.append({"kind": "word", "form": t, "lemma": lm, "hit": hits})
                for ki in hits:
                    key, path, wl = wls[ki]
                    a = per_key[ki].setdefault(("word", lm),
                        {"type": "word", "lemma": lm, "form": t, "pos": wl["pos"].get(lm, ""),
                         "freq": 0, "nsent": set(), "examples": []})
                    a["freq"] += 1; a["nsent"].add(si)
                    if len(a["examples"]) < 2: a["examples"].append(s)
                    hit_tok[ki].add((si, ti))
    return per_key, hit_tok, out_items, numbers_n, proper_n, len(out_items)

# ---------------- 归属分层（最低层只归一层） ----------------
def layering(out_items, wls_keys):
    layer_counts = collections.Counter()
    layer_meta = {}
    out_count = proper_count = num_count = 0
    for it in out_items:
        if it["kind"] == "num":
            num_count += 1
        elif it["kind"] == "proper":
            proper_count += 1
        elif it["kind"] == "out":
            out_count += 1
        else:
            layer = min(it["hit"], default=-1) if it["hit"] else -1
            key = it["lemma"]
            layer_counts[wls_keys[layer] if layer >= 0 else "out"] += 1
            meta = layer_meta.setdefault(key, {"layer": wls_keys[layer] if layer >= 0 else "out",
                                                "pos": "", "freq": 0, "form": it["form"]})
            meta["freq"] += 1
    return layer_counts, layer_meta, out_count, proper_count, num_count

# ---------------- N-gram + 互信息(MI) 候选短语提取 ----------------
# 学理依据: 固定词组是多词语块的强共现。用滑动窗口取 2~4 gram,
# 以内部相邻 bigram 的平均互信息 MI=log2(N*C(ab)/(C(a)C(b))) 衡量其键合强度,
# 频次(数量) + MI(键合质量) 双通道筛选, 交由教师核注后入 phrase_extra.json。
FUNC = set("""a an the of to in on at for with by from up down out off over under about into
through after before between against and or but so if because while when that this these those it its
we you they he she i my your our their his her is are was were be been being am have has had do does
did will would can could may might must shall should not no yes there here than as then also very
too just only more most some any all each every both other another 's 've 'll 're 'm 'd n't so""".split())

def extract_mwe_candidates(sents, wls, min_freq=2, nlim=4):
    """滑动窗口提取词群候选, MI 打分排序。返回 [{"phrase","len","freq","mi"}, ...]。"""
    known = set()
    for k, p, wl in wls:
        known |= set(wl["phrs"])
    unigram = collections.Counter(); bigram = collections.Counter()
    sents_toks = []
    for s in sents:
        toks = [t for t in (base_word(x).lower() for x in tokenize(s)) if t]
        sents_toks.append(toks); unigram.update(toks)
        for i in range(len(toks) - 1):
            bigram[(toks[i], toks[i + 1])] += 1
    total = sum(unigram.values())
    ngram = collections.Counter()
    for toks in sents_toks:
        n = len(toks)
        for L in range(2, min(nlim, n) + 1):
            for i in range(n - L + 1):
                g = tuple(toks[i:i + L])
                # 全功能词(无实词)不构成词组
                if not any(t not in FUNC for t in g):
                    continue
                ngram[g] += 1
    def bigram_mi(a, b):
        cab = bigram.get((a, b), 0); ca = unigram.get(a, 0); cb = unigram.get(b, 0)
        if not (cab and ca and cb): return -1e6
        return math.log2(total * cab / (ca * cb))
    out = []
    for g, f in ngram.items():
        if f < min_freq or " ".join(g) in known:
            continue
        mis = [bigram_mi(g[k - 1], g[k]) for k in range(1, len(g))] or [-1e6]
        out.append({"phrase": " ".join(g), "len": len(g), "freq": f, "mi": round(sum(mis) / len(mis), 2)})
    out.sort(key=lambda r: (-r["mi"], -r["freq"]))
    return out

# ---------------- 报告 ----------------
def write_report(wls, per_key, hit_tok, out_items, numbers_n, proper_n, total, cefr, out_stem, source,
                 mwe=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    thin = Side(style="thin", color="D9D9D9"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="2E8B57"); hdr_font = Font(color="FFFFFF", bold=True)
    hit_fill = PatternFill("solid", fgColor="E2F0E6")
    wb = openpyxl.Workbook()

    ws = wb.active; ws.title = "覆盖率"
    ws.append(["层级词表", "命中token", "token占比", "命中词条", "词表总词条"])
    for c in range(1, 6):
        cell = ws.cell(1, c); cell.fill = hdr_fill; cell.font = hdr_font
    keys = [k for k, _, _ in wls]
    for ki, (key, path, wl) in enumerate(wls):
        ws.append([key, len(hit_tok[ki]), "%.1f%%" % (100.0 * len(hit_tok[ki]) / total if total else 0),
                   len(per_key[ki]), len(wl["words"]) + len(wl["phrs"])])

    ws.append([])
    ws.append(["—— 归属分层（最低层只归一层，闭环） ——"])
    lc, lm, outc, propc, numc = layering(out_items, keys)
    ws.append(["层级", "token数", "占比"])
    for k in keys + ["out"]:
        ws.append([k, lc[k], "%.1f%%" % (100.0 * lc[k] / total if total else 0)])
    ws.append(["专名", propc, "%.2f%%" % (100.0 * propc / total if total else 0)])
    ws.append(["数字", numc, "%.2f%%" % (100.0 * numc / total if total else 0)])
    merged = len(set().union(*hit_tok)) if hit_tok else 0
    ws.append(["合并命中(任一表)", merged, "%.1f%%" % (100.0 * merged / total if total else 0), "-", "-"])
    ws.append(["全部token", total, "100.0%", "-", "-"])
    for i, w in enumerate((14, 11, 12, 10, 12), 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    COL = ["词/短语", "词性", "类型", "频次", "跨句", "CEFR", "原文例句"]; WID = (26, 16, 8, 8, 8, 8, 90)
    # 专名词表：阅读即识常见地名/机构名/物名（kind=proper 的词，零联网，作认识清单）
    prop = collections.Counter(it["form"] for it in out_items if it["kind"] == "proper")
    if prop:
        pn = wb.create_sheet("专名词")
        pn.append(["专名词表(认识清单)", "共 %d 个专名单词 · 地名/机构/物名等首字母大写词" % len(prop)])
        for _c in range(1, 3):
            cell = pn.cell(2, _c); cell.fill = hdr_fill; cell.font = hdr_font
        pn.append(["专名", "频次"])
        for w, n in prop.most_common():
            pn.append([w, n])
            for _c in range(1, 3):
                cell = pn.cell(pn.max_row, _c); cell.border = border; cell.fill = hit_fill
        pn.column_dimensions["A"].width = 26; pn.column_dimensions["B"].width = 8
    # KET 词表词头集，用于"PET 命中排除 KET 已含词"——只取真正在 KET 词表收录的 lemma
    _ket_words = next((wl["words"] for k, _, wl in wls if k == "KET"), frozenset())
    for ki, (key, path, wl) in enumerate(wls):
        heads = sorted(per_key[ki].values(), key=lambda a: (-a["freq"], a["lemma"]))
        sub = wb.create_sheet("命中-" + str(key))
        sub.append(["词表", "命中词条数", "共 %d 词条 · 来源:%s" % (len(heads), source)])
        for c in range(1, len(COL) + 1):
            cell = sub.cell(2, c); cell.fill = hdr_fill; cell.font = hdr_font
        sub.append(COL)
        for a in heads:
            sub.append([a["lemma"], a["pos"], "短语" if a["type"] == "phrase" else "单词",
                        a["freq"], len(a["nsent"]), cefr_level(cefr, a["lemma"]),
                        a["examples"][0] if a["examples"] else ""])
            for c in range(1, len(COL) + 1):
                cell = sub.cell(sub.max_row, c); cell.border = border; cell.fill = hit_fill
        for i, w in enumerate(WID, 1):
            sub.column_dimensions[get_column_letter(i)].width = w
        sub.freeze_panes = "A4"
        # 对 PET 层：另出一张排除 KET 已含词的"新增"表（真正 B1 新授词，而非 KET 已会词）
        if key == "PET":
            sub2 = wb.create_sheet("命中-PET(排除KET)")
            sub2.append(["词表", "PET新增(B1)词条数",
                         "共 %d 词条 · 已在KET词表中出现的词已剔除 · 来源:%s"
                         % (len([a for a in heads if a["lemma"] not in _ket_words]), source)])
            for c in range(1, len(COL) + 1):
                cell = sub2.cell(2, c); cell.fill = hdr_fill; cell.font = hdr_font
            sub2.append(COL)
            for a in heads:
                if a["lemma"] in _ket_words:
                    continue
                sub2.append([a["lemma"], a["pos"], "短语" if a["type"] == "phrase" else "单词",
                             a["freq"], len(a["nsent"]), cefr_level(cefr, a["lemma"]),
                             a["examples"][0] if a["examples"] else ""])
                for c in range(1, len(COL) + 1):
                    cell = sub2.cell(sub2.max_row, c); cell.border = border; cell.fill = hit_fill
            for i, w in enumerate(WID, 1):
                sub2.column_dimensions[get_column_letter(i)].width = w
            sub2.freeze_panes = "A4"

    oc = collections.Counter(it["lemma"] for it in out_items if it["kind"] == "out")
    ws_o = wb.create_sheet("超纲词")
    # 超纲词原句：同一词收集出现过的原句（去重、去空），供分支做"原文照搬"。
    out_sent = {}
    for it in out_items:
        if it["kind"] == "out" and it.get("sent"):
            out_sent.setdefault(it["lemma"], it["sent"])
    ws_o.append(["词", "频次", "CEFR", "原文例句"])
    for c in range(1, 5):
        cell = ws_o.cell(1, c); cell.fill = hdr_fill; cell.font = hdr_font
    for w, c in oc.most_common():
        ws_o.append([w, c, cefr_level(cefr, w), out_sent.get(w, "")])
    for i, w in enumerate((24, 8, 8, 60), 1):
        ws_o.column_dimensions[get_column_letter(i)].width = w
    # 候选待核短语: N-gram + MI 提取, 供教师核对后入 phrase_extra.json
    if mwe:
        ws_m = wb.create_sheet("候选待核短语")
        ws_m.append(["待核短语", "词数", "频次", "互信息MI", "核注后请加入 wordlists/phrase_extra.json"])
        for c in range(1, 6):
            cell = ws_m.cell(1, c); cell.fill = hdr_fill; cell.font = hdr_font
        for r in mwe:
            ws_m.append([r["phrase"], r["len"], r["freq"], r["mi"], ""])
            for c in range(1, 6):
                cell = ws_m.cell(ws_m.max_row, c); cell.border = border; cell.fill = hit_fill
        for i, w in enumerate((28, 6, 8, 12, 40), 1):
            ws_m.column_dimensions[get_column_letter(i)].width = w
        ws_m.freeze_panes = "A2"
    wb.save(out_stem + ".xlsx")

    keys_ = [k for k, _, _ in wls]
    for it in out_items:
        if it["kind"] == "word":
            it["layer"] = keys_[min(it["hit"])] if it["hit"] else "out"
            it["cefr"] = cefr_level(cefr, it["lemma"])
    out = {"source": source, "total": total, "numbers": numbers_n, "proper": proper_n,
           "tokens": out_items}
    with open(out_stem + ".json", "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    return out_stem + ".xlsx"

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--text", default=None)
    ap.add_argument("--file", default=None)
    ap.add_argument("--xlsx", default=None)
    ap.add_argument("--sheet", default="C4 句库时间轴")
    ap.add_argument("--col", type=int, default=2)
    ap.add_argument("--stdin", action="store_true")
    ap.add_argument("--out", default="vocab_report")
    ap.add_argument("--cefr", default=None, help="可选的 cefr.json 路径，用于附加难度标注")
    ap.add_argument("--db", default=None, help="词汇库SQLite路径，传则把本报告UPSERT入库")
    ap.add_argument("--doc", default=None, help="入库时的文章标识(默认用输入路径/时间戳)")
    ap.add_argument("--db-export", default=None, help="词汇库导出Excel视图的目标路径")
    ap.add_argument("--mwe", action="store_true",
                    help="N-gram+互信息提取候选短语, 输出'候选待核短语'sheet供教师核注入 phrase_extra.json")
    a = ap.parse_args()
    skill_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    # worktree 直接运行/部署均可用：优先脚本两级上级(部署位), 退化到本文件同级与当前工作目录
    for cand in (skill_dir, os.path.dirname(os.path.abspath(__file__)), os.getcwd()):
        if os.path.isdir(os.path.join(cand, "wordlists")):
            skill_dir = cand
            break
    wls = discover_wordlists(skill_dir)
    if not wls:
        sys.exit("未找到任何词表: %s/wordlists" % skill_dir)
    sents = load_sentences(a)
    if not sents:
        sys.exit("课文为空, 请提供课文输入")
    per_key, hit_tok, out_items, numbers_n, proper_n, total = assess(sents, wls)
    # cefr 来源: 显式 --cefr 优先; 否则自动加载 <skill>/wordlists/cefr.json(存在才用, 缺省留空不猜)
    cefr_path = a.cefr or os.path.join(skill_dir, "wordlists", "cefr.json")
    cefr = load_cefr(cefr_path if os.path.exists(cefr_path) else a.cefr)
    out_stem = os.path.join(os.getcwd(), a.out)
    mwe = extract_mwe_candidates(sents, wls) if a.mwe else None
    xlsx = write_report(wls, per_key, hit_tok, out_items, numbers_n, proper_n, total, cefr,
                        out_stem, " | ".join(k for k, _, _ in wls), mwe)
    if a.db:
        wb_ = _load_word_bank()
        _keys = [k for k, _, _ in wls]
        _src = " | ".join(_keys)
        # 汇总词项级记录: 合并 per_key(pos/例句) 与 out_items(layer/lemma)
        _items = []
        _seen = set()
        for it in out_items:
            if it["kind"] not in ("word", "out") or not it.get("lemma") or it["lemma"] in _seen:
                continue
            _seen.add(it["lemma"])
            _lm = it["lemma"]
            _ki = min(it["hit"]) if it["hit"] else -1
            _rec = per_key[_ki].get(("word", _lm)) if _ki >= 0 else None
            _items.append({"kind": "word",
                           "lemma": _lm,
                           "pos": (_rec["pos"] if _rec else ""),
                           "layer": _keys[_ki] if _ki >= 0 else "out",
                           "cefr": cefr_level(cefr, _lm),
                           "examples": (_rec["examples"][:2] if _rec and _rec.get("examples") else [])})
        if a.doc:
            doc_id = a.doc
        elif a.file:
            doc_id = os.path.abspath(a.file)
        else:
            import time; doc_id = "inline-" + str(int(time.time()))
        _n, _ins = wb_.upsert_report(a.db, _items, _src, doc_id)
        print("词库入库: 词数=%d 新增=%d -> %s" % (_n, _ins, a.db))
    if a.db_export:
        wb_ = _load_word_bank()
        wb_.export_views(a.db or "vocab.db", a.db_export)
        print("词库视图已导出:", a.db_export)
    print("saved:", xlsx)
    keys = [k for k, _, _ in wls]
    print("各层命中:", " ".join("%s=%.1f%%" % (k, 100.0 * len(h) / total if total else 0)
                              for (k, p, wl), h in zip(wls, hit_tok)))

if __name__ == "__main__":
    main()