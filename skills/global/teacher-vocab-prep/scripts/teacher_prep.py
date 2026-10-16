# -*- coding: utf-8 -*-
"""教师备课软分支·后处理
主技能 article-vocab-analyzer 的任务结果(分层+词性+频次+跨句Range+CEFR+超纲) → 加工成教师备课素材
(分层教学词卡 / 考点命题矩阵 / 超纲进阶补卡 / 单元考点统计)。

本分支只复用主技能的**静态分析结果**，不联网、不需 nltk，快且稳，适合备课。

用法:
  python teacher_prep.py --file 课文.txt --out 教师备课本
  python teacher_prep.py --text "..." --out 备课
  python teacher_prep.py --file 课文.txt --band "全员(基础)"   # 只挑某教学梯度
"""
import os, sys, re, json, io, subprocess, argparse

def _main_dir():
    """向上查找真实主技能目录(article-vocab-analyzer)，兼容多布局(平铺/skills、skills/global、skills/project)。"""
    d = os.path.dirname(os.path.abspath(__file__))
    subs = (("skills", "article-vocab-analyzer"), ("skills", "global", "article-vocab-analyzer"),
            ("skills", "project", "article-vocab-analyzer"), ("article-vocab-analyzer",))
    for _ in range(8):
        for sub in subs:
            cand = os.path.join(d, *sub)
            if os.path.exists(os.path.join(cand, "scripts", "assess_vocab.py")):
                return os.path.abspath(cand)
        d = os.path.dirname(d)
    return None
MAIN = _main_dir() or os.path.abspath("article-vocab-analyzer")
ASSESS = os.path.join(MAIN, "scripts", "assess_vocab.py")
# worktree 平铺优先：本文件同级有 assess_vocab.py 时用它(开发版最新)，否则回退部署位
_flat = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assess_vocab.py")
if os.path.exists(_flat):
    ASSESS = _flat

# ---------- 1. 调主技能：json(lemma/layer/cefr/forms) + report.xlsx(词性/频次/跨句/原句) ----------
def _load_assess():
    import importlib.util
    spec = importlib.util.spec_from_file_location("_av", ASSESS)
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
    return m

def run_main(text=None, file=None, stem="tp_tmp"):
    cmd = [sys.executable, ASSESS, "--out", stem]
    if text:
        cmd += ["--text", text]
    else:
        cmd += ["--file", file]
    subprocess.check_call(cmd)
    toks = json.load(io.open(stem + ".json", encoding="utf-8"))["tokens"]
    av = _load_assess()
    words = {}
    for t in toks:
        if t.get("kind") not in ("word", "out"):
            continue
        lm = t.get("lemma", "")
        if not lm:
            continue
        rec = words.setdefault(lm, {"layer": t.get("layer", ""), "cefr": t.get("cefr", ""), "forms": []})
        f = t.get("form", "")
        if f and f not in rec["forms"]:
            rec["forms"].append(f)
    # 显示口径（2026-10-13 定案，与词汇诊断页一致）：词原本身在课文出现 → 词原；
    # 否则该词族以不规则表层形式出现（came/went/children…）→ 显示表层形式。
    for lm, rec in words.items():
        bases = {av.base_word(f).lower() for f in rec["forms"]}
        rec["disp"] = lm if lm in bases else (av.display_word(rec["forms"][0], lm) if rec["forms"] else lm)
    try:
        import openpyxl
    except ImportError:
        return words, {}
    wb = openpyxl.load_workbook(stem + ".xlsx", read_only=True)
    meta = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        for row in ws.iter_rows(values_only=True):
            if not row or not row[0]:
                continue
            c = list(row)
            if sn.startswith("命中"):   # 词/短语|词性|类型|频次|跨句|CEFR|原文例句
                w = c[0]
                if w and isinstance(w, str) and len(w) < 40 and w not in meta:
                    meta[w] = {"pos": c[1] or "", "type": c[2] or "", "freq": c[3], "range": c[4],
                               "cefr": c[5] or "", "example": c[6] or ""}
            elif sn == "超纲词":        # 词|频次|CEFR|原文例句
                w = c[0]
                meta.setdefault(w, {}).update({"freq": c[1], "cefr": c[2] or "", "pos": "",
                                               "example": c[3] if len(c) > 3 else ""})
    wb.close()
    return words, meta

# ---------- 2. 教学梯度分层 + 词性集合 ----------
BAND = {"KET": "全员(基础)", "PET": "全员(基础)", "NGSL": "中阶", "NAWL": "中强", "out": "中强(超纲/进阶)"}
_POSOK = {"n", "v", "adj", "adv", "prep", "pron", "det", "conj", "num", "exclam", "modal"}
def _base(p):
    p = re.sub(r"[^a-z]", "", p.lower())
    if p in ("vt", "vi", "vti", "intransitive", "transitive", "verb"):
        return "v"
    if p.startswith("adj") or p in ("p", "modify"):
        return "adj"
    if p.startswith("adv"):
        return "adv"
    if p.startswith("n"):
        return "n"
    if p.startswith("prep"):
        return "prep"
    if p.startswith("art") or p.startswith("det"):
        return "det"
    if p.startswith("pron"):
        return "pron"
    if p.startswith("conj"):
        return "conj"
    if p.startswith("num"):
        return "num"
    if p.startswith("exclam") or p.startswith("excl") or p.startswith("interj"):
        return "exclam"
    if p.startswith("modal") or p.startswith("aux"):
        return "modal"
    return p
def posset(pos):
    if not pos or pos.strip().upper() in ("N/A", "NA", "-"):
        return set()
    s = pos.replace("&", " ").replace("/", " ").replace(",", " ")
    return {b for b in (_base(x) for x in re.split(r"\s+", s)) if b in _POSOK}

def _band(layer):
    return BAND.get(layer, layer or "")

# ---------- 3. 命题参考(自动化建议, 须人工复核) ----------
def _quiz_tags(pos):
    if not pos:
        return ["挖空辨义；词义/搭配选填"]
    tags = []
    if "v" in pos:
        tags.append("时态/语态/词形")
    if "n" in pos:
        tags.append("单复数/冠词搭配")
    if "adj" in pos:
        tags.append("比较级/最高级/搭配")
    if "adv" in pos or "prep" in pos:
        tags.append("搭配/句内功能")
    if not tags:
        tags = ["挖空辨义"]
    return tags[:3]

def _blanks(example, forms):
    """把原句中首个目标词(含其屈折形式)替换为 ____ 作挖空位；无原句/词未现返回 ''。"""
    if not example or not forms:
        return ""
    import re as _re
    for w in forms:
        m = _re.search(r"(?i)\b" + _re.escape(w) + r"\b", example)
        if m:
            return (example[:m.start()] + "____" + example[m.end():]).strip()
    return ""

# ---------- 4. 组装教师备课本 ----------
def main():
    ap = argparse.ArgumentParser(description="教师备课(主技能静态结果 → 分层词卡/考点矩阵/超纲补卡)")
    ap.add_argument("--file", default=None)
    ap.add_argument("--text", default=None)
    ap.add_argument("--out", default="教师备课本")
    ap.add_argument("--band", default=None, help="只保留某教学梯度(逗号分隔), 如 全员(基础),中阶")
    a = ap.parse_args()
    if a.text is None and a.file is None:
        ap.error("需要 --file 或 --text")

    words, meta = run_main(a.text, a.file, "__tp_" + a.out)
    import openpyxl

    rows = []
    for lm in words:
        w = words[lm]
        m = meta.get(lm, {})
        fr = m.get("freq")
        freq = int(fr) if isinstance(fr, (int, float)) else 0
        rows.append([
            _band(w["layer"]), w["layer"], w.get("disp", lm), m.get("pos", ""), freq,
            m.get("range") or "", w["cefr"] or m.get("cefr", "") or "", m.get("example", ""),
            m.get("type", ""),
        ])
    disp2lm = {w.get("disp", lm): lm for lm, w in words.items()}   # 显示词 → 词原（词族回查用）
    # 分层词卡：超纲排在末；其余按频次降序
    rows.sort(key=lambda r: (1 if r[1] == "out" else 0, -(r[4] if r[4] else 0)))
    if a.band:
        want = set(x.strip() for x in a.band.split(","))
        rows = [r for r in rows if r[0] in want]

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "分层词卡"
    ws.append(["教学梯度", "归属层", "词汇", "词性", "频次", "跨句Range", "CEFR", "原文照搬", "类型"])
    for r in rows:
        ws.append(r)

    # 考点命题矩阵：只对实词(n/v/adj/adv)给命题参考 + 挖空位
    ws2 = wb.create_sheet("考点命题矩阵")
    ws2.append(["教学梯度", "归属层", "词汇", "词性", "CEFR", "原文照搬", "可挖空(填空)", "命题参考(建议)", "考情提示"])
    for r in rows:
        disp, pos, cefr = r[2], r[3], r[6]
        tot = posset(pos)
        if not (tot & {"n", "v", "adj", "adv"}):
            continue
        src = disp2lm.get(disp, disp)
        hint = "基础必考(A1/B1)" if r[1] in ("KET", "PET") or (cefr and cefr in ("A1", "A2", "B1")) else "进阶讲读(CET可萃取)"
        ws2.append([r[0], r[1], disp, pos, cefr, r[7], _blanks(r[7], [src] + words[src].get("forms", [])),
                    "；".join(_quiz_tags(tot)), hint])

    # 超纲/进阶补卡：out + NAWL + NGSL(中阶) 按频次降序
    ws3 = wb.create_sheet("超纲进阶补卡")
    ws3.append(["教学梯度", "归属层", "词汇", "频次", "CEFR", "原文照搬"])
    adv = [r for r in rows if r[1] in ("out", "NAWL", "NGSL")]
    adv.sort(key=lambda r: -(r[4] if isinstance(r[4], int) else 0))
    for r in adv:
        ws3.append([r[0], r[1], r[2], r[4], r[6], r[7]])

    # 单元考点统计：层词数 + 占比 + 高频词top
    from collections import Counter
    ws4 = wb.create_sheet("单元考点统计")
    lay = Counter(r[1] or "未知" for r in rows)
    ws4.append(["归属层", "词条数", "教学梯度"])
    for k in [x for x in ("A1_Starters", "A1_Movers", "A2_Flyers", "KET", "PET", "FCE", "CAE",
                          "NGSL", "NAWL", "out") if x in lay] + sorted(set(lay) - {"A1_Starters", "A1_Movers", "A2_Flyers", "KET", "PET", "FCE", "CAE", "NGSL", "NAWL", "out"}):
        if k in lay:
            ws4.append([k, lay[k], _band(k)])
    total = sum(lay.values()) or 1
    ws4.append([]); ws4.append(["占比(词条口径)", ])
    for k, v in lay.items():
        ws4.append(["", "%s: %.1f%%" % (k, v * 100.0 / total)])
    ws4.append([]); ws4.append(["高频词TOP15(按频次)", "频次", "层", "教学梯度"])
    for r in rows[:15]:
        ws4.append([r[2], r[4], r[1], r[0]])

    wb.save(a.out + ".xlsx")
    for ext in (".json", ".xlsx"):
        try:
            os.remove("__tp_" + a.out + ext)
        except OSError:
            pass
    print("完成: %d 词 -> %s.xlsx (分层词卡/考点命题矩阵/超纲进阶补卡/单元考点统计)" % (len(rows), a.out))

if __name__ == "__main__":
    main()