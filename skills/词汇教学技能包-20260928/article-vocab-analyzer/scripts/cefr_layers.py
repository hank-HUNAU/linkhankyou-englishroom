# -*- coding: utf-8 -*-
"""CEFR 多源层：把「单点 CEFR 标签」升级为「多源 + 来源留痕 + 冲突双标注」。

设计约定（2026-10-20 定案）
--------------------------
1. 取值顺序（主标注 level 取第一个有值的来源）：
       oxford(Oxford3000/5000) → cefrj(CEFR-J v1.6) → octanove(C1/C2) → efllex(峰值)
   Oxford 存在时即为 Oxford，**不被后来者覆盖**。
2. 冲突只在 {oxford, cefrj, octanove} 之间判定——三者同为「等级赋值」构念。
   EFLLex 是「等级频率分布」资源，构念不同：只作补缺，**不参与冲突判定**。
3. 冲突词条保留全部来源等级，主标注仍为 Oxford，禁止静默覆盖。
4. 所有来源均需署名，见 CITATIONS；第三方数据不随本技能分发，用
   scripts/fetch_cefr_sources.py 按需下载到 wordlists/cefr_sources/。

用法：
    from cefr_layers import load_sources, build, lookup
    srcs = load_sources(cache_dir)          # 缺源自动降级，不报错
    ext  = build(oxford_words, srcs)        # 合并后的词条表
    level, s, alt, conflict = lookup(ext, "abandon")
"""
import os, json, csv, re

ORDER = ["A1", "A2", "B1", "B2", "C1", "C2"]
GRADED = ("oxford", "cefrj", "octanove")   # 参与冲突判定的「等级赋值」源
ALL_SOURCES = ("oxford", "cefrj", "octanove", "efllex")

SOURCES = {
    "oxford": {"name": "Oxford 3000/5000", "levels": "A1–C1",
               "license": "词典版权，无官方批量下载；技能内置 cefr.json 为已落地版本",
               "need": "技能自带（wordlists/cefr.json）"},
    "cefrj": {"name": "CEFR-J Wordlist v1.6（TUFS 投野研究室）", "levels": "A1–B2",
              "license": "适当引用前提下，研究/教育/商用均免费；改编后仍须引用",
              "need": "CEFRJ_wordlist_ver1.6.zip（fetch 脚本下载）"},
    "octanove": {"name": "Octanove Vocabulary Profile C1/C2 v1.0", "levels": "C1–C2",
                 "license": "CC BY-SA 4.0",
                 "need": "octanove-vocabulary-profile-c1c2-1.0.csv"},
    "efllex": {"name": "EFLLex（UCLouvain / CENTAL）", "levels": "A1–C1",
               "license": "研究/教学免费，须引用论文",
               "need": "EFLLex.tsv"},
}

CITATIONS = [
    "Oxford 3000 / Oxford 5000 — Oxford University Press（Oxford Learner's Dictionaries）；"
    "CEFR 星标用于难度坐标，无官方批量下载。",
    "The CEFR-J Wordlist Version 1.6. Compiled by Yukio Tono, Tokyo University of Foreign Studies. "
    "Retrieved from http://www.cefr-j.org/download.html on dd/mm/yy.",
    "Octanove Vocabulary Profile C1/C2 (ver 1.0), Octanove Labs — CC BY-SA 4.0 "
    "(https://creativecommons.org/licenses/by-sa/4.0/).",
    "Dürlich, L. & François, T. (2018). EFLLex: A Graded Lexical Resource for Learners of English "
    "as a Foreign Language. LREC 2018. DOI: 10.63317/43e9y9se8gyj",
]

TOKEN_RE = re.compile(r"[a-z][a-z'\-]*")


def _read_csv_levels(path, level_col="CEFR", head_col="headword", pos_col="pos"):
    out = {}
    with open(path, encoding="utf-8", errors="replace") as f:
        for row in csv.DictReader(f):
            h = (row.get(head_col) or "").strip().lower()
            lv = (row.get(level_col) or "").strip().upper()
            if h and lv in ORDER:
                out.setdefault(h, {"level": lv, "pos": (row.get(pos_col) or "").strip()})
    return out


def _read_xlsx_cefrj(path, sheet="ALL_sep"):
    """CEFR-J v1.6 Excel：ALL_sep 表已把斜杠词头拆开，便于查表"""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sn = sheet if sheet in wb.sheetnames else wb.sheetnames[1]
    out = {}
    rows = wb[sn].iter_rows(values_only=True)
    next(rows, None)
    for r in rows:
        if not r or not r[0]:
            continue
        h = str(r[0]).strip().lower()
        lv = (str(r[2]).strip().upper() if len(r) > 2 and r[2] else "")
        if h and lv in ORDER:
            out.setdefault(h, {"level": lv, "pos": (str(r[1]).strip() if len(r) > 1 and r[1] else "")})
    return out


def _read_efllex(path):
    """EFLLex.tsv：按等级频率取峰值等级（A1–C1）"""
    out = {}
    with open(path, encoding="utf-8", errors="replace") as f:
        header = f.readline().rstrip("\n").split("\t")
        cols = {}
        for i, c in enumerate(header):
            m = re.match(r"level_freq@(a1|a2|b1|b2|c1)$", c)
            if m:
                cols[m.group(1)] = i
        if not cols:
            return out
        for line in f:
            p = line.rstrip("\n").split("\t")
            if len(p) < 8:
                continue
            w = p[0].strip().lower()
            if not w:
                continue
            vec = {}
            for lv, i in cols.items():
                try:
                    vec[lv] = float(p[i])
                except (ValueError, IndexError):
                    vec[lv] = 0.0
            if not any(vec.values()):
                continue
            peak = max(vec, key=lambda k: vec[k]).upper()
            out.setdefault(w, {"level": peak, "pos": p[1] if len(p) > 1 else "",
                               "vec": {k: round(v, 2) for k, v in vec.items()}})
    return out


def load_sources(cache_dir, oxford_words=None):
    """从缓存目录加载四个源；缺哪个源就降级（不报错）。oxford_words 由调用方传入。"""
    srcs = {}
    if oxford_words:
        srcs["oxford"] = {w: {"level": v, "pos": ""} for w, v in oxford_words.items()}
    for fn, key, reader in (
        ("octanove-vocabulary-profile-c1c2-1.0.csv", "octanove", _read_csv_levels),
        ("EFLLex.tsv", "efllex", _read_efllex),
    ):
        p = os.path.join(cache_dir, fn)
        if os.path.exists(p):
            srcs[key] = reader(p)
    cefrj_dir = os.path.join(cache_dir, "cefrj16")
    if os.path.isdir(cefrj_dir):
        for fn in os.listdir(cefrj_dir):
            if fn.lower().endswith(".xlsx"):
                srcs["cefrj"] = _read_xlsx_cefrj(os.path.join(cefrj_dir, fn))
                break
    return srcs


def build(oxford_words=None, srcs=None):
    """合并成词条表：{word: {level, src, pos, oxford, cefrj, octanove, efllex, conflict}}"""
    srcs = srcs or {}
    entries = {}
    words = set()
    for k in ALL_SOURCES:
        words |= set(srcs.get(k, {}))
    for w in words:
        vals = {}
        for k in ALL_SOURCES:
            v = srcs.get(k, {}).get(w)
            if v:
                vals[k] = v["level"] if isinstance(v, dict) else v
        if not vals:
            continue
        conflict = len({vals[k] for k in GRADED if k in vals}) > 1
        for k in ("oxford", "cefrj", "octanove", "efllex"):
            if k in vals:
                level, s = vals[k], k
                break
        pos = ""
        for k in ("cefrj", "octanove", "efllex"):
            d = srcs.get(k, {}).get(w)
            if isinstance(d, dict) and d.get("pos"):
                pos = d["pos"]
                break
        entries[w] = {"level": level, "src": s, "pos": pos,
                      "oxford": vals.get("oxford", ""), "cefrj": vals.get("cefrj", ""),
                      "octanove": vals.get("octanove", ""), "efllex": vals.get("efllex", ""),
                      "conflict": conflict}
    return entries


def lookup(entries, lemma, extra_candidates=None):
    """查合并表。extra_candidates 可传入额外的词形候选（如 -ed/-ing 修正候选）。
    返回 (level, src, alt, conflict)；alt 仅在冲突时非空。"""
    cands = [lemma] + list(extra_candidates or [])
    seen = set()
    for c in cands:
        if not c or c in seen:
            continue
        seen.add(c)
        e = entries.get(c)
        if e:
            alt = {k: e[k] for k in ALL_SOURCES if e.get(k)} if e.get("conflict") else {}
            return e["level"], e["src"], alt, bool(e.get("conflict"))
    return "", "", {}, False


def summary(entries):
    """来源/冲突统计，供报告页脚与 sheet 使用"""
    import collections
    src = collections.Counter(e["src"] for e in entries.values())
    lv = collections.Counter(e["level"] for e in entries.values())
    return {"sources": dict(src), "levels": dict(lv),
            "conflicts": sum(1 for e in entries.values() if e["conflict"]),
            "total": len(entries)}
